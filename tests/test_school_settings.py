"""Configuration/workflow integration, atomic failures and tenant isolation."""
import os,re,unittest
from datetime import date
os.environ['DATABASE_URL']='sqlite:///:memory:'
from app import app,db,Student,Employee
from school_settings import EXAM_PAGES,HRM_PAGES

class SchoolSettingsTests(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True,IDENTITY_ENFORCE_IN_TESTS=False,SECRET_KEY='settings-test')
        self.ctx=app.app_context();self.ctx.push();db.drop_all();db.create_all()
        self.m=app.extensions['school_settings_models'];self.rm=app.extensions['record_models'];self.ops=app.extensions['ops_models']
        for cat,vals in {'campus':['Campus A','Campus B'],'class_name':['Class A'],'section':['A','B'],'designation_group':['Teaching','Office'],'designation':['Class Teacher','Clerk'],'department':['Academics'],'employment_type':['Permanent'],'marital_status':['Married']}.items():
            for v in vals:db.session.add(self.rm['settings'](category=cat,value=v,class_name=''))
        self.e=Employee(employee_no='E-1',name='Fictional Teacher',role='Class Teacher',status='Active')
        self.clerk=Employee(employee_no='E-2',name='Fictional Clerk',role='Clerk',status='Active')
        self.s=Student(admission_no='S-1',name='Fictional Pupil',class_name='Class A',status='Active')
        db.session.add_all([self.e,self.clerk,self.s]);db.session.flush()
        for e in [self.e,self.clerk]:db.session.add(self.rm['employee'](employee_id=e.id,data={'campus':'Campus A','joining_date':'2000-01-01','email':'retained@example.invalid'}))
        db.session.add(self.rm['student'](student_id=self.s.id,data={'campus':'Campus A','section':'A'}));db.session.commit()
        self.c=app.test_client();self.c.get('/workspace/core/hrm')
        with self.c.session_transaction() as session:self.csrf=session['record_csrf']
    def tearDown(self):db.session.remove();db.drop_all();self.ctx.pop()
    def post(self,module,page,**values):
        response=self.c.post('/workspace/core/'+module+'/'+page,data={'_csrf':self.csrf,'version':'0',**values})
        if response.status_code==302:self.assertEqual(self.c.get(response.location).status_code,200,response.location)
        return response
    def config(self,key):return db.session.get(self.m['setting'],key)
    def item(self,kind):return self.m['item'].query.filter_by(kind=kind).one()
    def shift(self,**kw):
        data=dict(name='Morning',direction='in',designations=['Class Teacher'],**{'time_'+str(i):'07:45' for i in range(7)});data.update(kw)
        return self.post('hrm','shift-templates',**data)
    def test_every_settings_page_read_only_renders_unique_ids(self):
        for module,pages in [('exam',EXAM_PAGES),('hrm',HRM_PAGES)]:
            for key,_ in pages:
                page=self.c.get(f'/workspace/core/{module}/{key}');self.assertEqual(page.status_code,200,key)
                ids=re.findall(rb'\bid="([^"]+)"',page.data);self.assertEqual(len(ids),len(set(ids)),key)
        self.assertEqual(self.m['item'].query.count(),0);self.assertEqual(self.m['setting'].query.count(),0)
    def test_csrf_stale_version_unknown_action_and_validation_atomic(self):
        path='/workspace/core/exam/display'
        self.assertEqual(self.c.post(path,data={'decimals':'2'}).status_code,400)
        self.assertEqual(self.post('exam','display',decimals='NaN').status_code,422)
        self.assertEqual(self.post('exam','display',decimals='2',written='Theory').status_code,302)
        self.assertEqual(self.post('exam','display',decimals='1').status_code,422)
        self.assertEqual(self.post('exam','display',version='1',action='bogus').status_code,422)
        self.assertEqual(self.config('exam:display').data['written'],'Theory')
    def test_daily_lock_blocks_direct_posts_unlocks_only_last_day(self):
        self.assertEqual(self.post('hrm','attendance-lock',daily='2026-09-10',monthly='2026-08').status_code,302)
        values={'_csrf':self.csrf,'day':'2026-09-10','status_'+str(self.e.id):'Present'}
        self.assertEqual(self.c.post('/workspace/attendance/staff',data=values).status_code,422)
        self.assertEqual(self.ops['daily'].query.count(),0)
        self.assertEqual(self.post('hrm','attendance-lock',version='1',action='unlock-daily',daily='2099-01-01').status_code,302)
        self.assertEqual(self.config('hrm:attendance-lock').data['daily'],'2026-09-09')
        self.assertEqual(self.c.post('/workspace/attendance/staff',data=values).status_code,302)
        values['day']='2026-09-09';self.assertEqual(self.c.post('/workspace/attendance/staff',data=values).status_code,422)
        values['day']='2026-09-08';values['status_'+str(self.s.id)]='Present'
        self.assertEqual(self.c.post('/workspace/attendance/student',data=values).status_code,302)
    def test_monthly_lock_applies_to_whole_month_daily_still_applies(self):
        self.post('hrm','attendance-lock',daily='2026-09-03',monthly='2026-08')
        path='/workspace/core/hrm/monthly-attendance';base={'_csrf':self.csrf,'employee_id':str(self.e.id),'month':'2026-08','status_2026-08-20':'Present'}
        self.assertEqual(self.c.post(path,data=base).status_code,422)
        self.assertEqual(self.post('hrm','attendance-lock',version='1',action='unlock-monthly').status_code,302)
        self.assertEqual(self.config('hrm:attendance-lock').data['monthly'],'2026-07')
        # The independent daily lock still blocks this old month.
        self.assertEqual(self.c.post(path,data=base).status_code,422)
        self.assertEqual(self.post('hrm','attendance-lock',version='2',daily='',monthly='2026-07').status_code,302)
        self.assertEqual(self.c.post(path,data=base).status_code,302)
        self.assertEqual(self.ops['daily'].query.one().status,'Present')
    def test_hidden_employee_fields_preserve_saved_values_and_cannot_be_forged(self):
        self.assertEqual(self.post('hrm','employee-fields',visible_notes='1',visible_photo='1',visible_documents='1').status_code,302)
        form=self.c.get('/records/employee/1/edit').data
        self.assertNotIn(b'name="email"',form);self.assertIn(b'name="name"',form)
        employee=dict(_csrf=self.csrf,name='Updated Fictional Teacher',father_name='Guardian',gender='Female',marital_status='Married',phone='00000000000',campus='Campus A',department='Academics',designation_group='Teaching',role='Class Teacher',employment_type='Permanent',employee_no='E-1',joining_date='2000-01-01',email='forged@example.invalid',save_action='back')
        self.assertEqual(self.c.post('/records/employee/1/edit',data=employee).status_code,302)
        detail=self.rm['employee'].query.filter_by(employee_id=1).one();self.assertEqual(detail.data['email'],'retained@example.invalid')
        self.assertEqual(self.e.name,'Updated Fictional Teacher')
        self.assertEqual(self.post('hrm','employee-fields',version='1',required_email='1').status_code,422)
        self.assertEqual(self.post('hrm','employee-fields',version='1',action='reset').status_code,302)
        self.assertIn(b'name="email"',self.c.get('/records/employee/1/edit').data)
    def test_shift_applicability_assignment_and_in_use_delete(self):
        self.assertEqual(self.shift().status_code,302);r=self.item('shift-template')
        path='/workspace/core/hrm/assign-shifts'
        values={'_csrf':self.csrf,'version':'0','action':'save','employee_ids':str(self.clerk.id),f'in_{self.clerk.id}_0':str(r.id)}
        self.assertEqual(self.c.post(path,data=values).status_code,422);self.assertIsNone(self.config('hrm:shift-assignment'))
        values.update(employee_ids=str(self.e.id));values[f'in_{self.e.id}_0']=str(r.id)
        self.assertEqual(self.c.post(path,data=values).status_code,302)
        page=self.c.get('/workspace/attendance/staff?day=2026-09-07');self.assertIn(b'07:45',page.data)
        self.assertEqual(self.post('hrm','shift-templates',id=r.id,version='1',action='delete').status_code,422)
        self.assertEqual(self.c.get(path+'?unassigned=1').status_code,200)
    def test_student_section_shift_overrides_default(self):
        def create(name,section,in_time):
            values=dict(name=name,campus='Campus A',classes=['Class A'],section=section,schedule_type='combined')
            values.update({'in_'+str(i):in_time for i in range(7)});values.update({'out_'+str(i):'13:00' for i in range(7)})
            return self.post('exam','student-shifts',**values)
        self.assertEqual(create('Default','','08:00').status_code,302);self.assertEqual(create('Section A','A','09:00').status_code,302)
        self.assertIn(b'09:00',self.c.get('/workspace/attendance/student?day=2026-09-07').data)
        self.assertEqual(create('Duplicate','A','10:00').status_code,422)
    def test_subject_assignment_rejects_unlinked_award_subject(self):
        self.post('exam','subject-groups',name='Academics')
        self.assertEqual(self.post('exam','subjects',name='Math',group='Academics',diary='on').status_code,302)
        self.assertEqual(self.post('exam','class-subjects',name='Class A',campus='Campus A',subjects=['Math']).status_code,302)
        form=dict(_csrf=self.csrf,name='Bad List',campus='Campus A',class_name='Class A',subject=['Other'],exam_date=['2026-09-07'],test_type=['Written'],total=['30'])
        self.assertEqual(self.c.post('/workspace/examination/awards',data=form).status_code,422)
        self.assertEqual(self.ops['workflow'].query.filter_by(kind='awards').count(),0)
        form.update(name='Good List',subject=['Math']);self.assertEqual(self.c.post('/workspace/examination/awards',data=form).status_code,302)
        subject=self.item('subject');self.assertEqual(self.post('exam','subjects',id=subject.id,version='1',action='delete').status_code,422)
    def test_grades_use_unrounded_percentage_in_results(self):
        self.assertEqual(self.post('exam','grading',name='School policy',minimum=['0','80'],maximum=['79.999999','100'],grade=['B','A'],remarks=['Good','Excellent']).status_code,302)
        policy=self.item('grading');self.assertEqual(self.post('exam','class-grading',name='Class policy',campus='Campus A',classes=['Class A'],policy=str(policy.id)).status_code,302)
        award=self.ops['workflow'](kind='awards',status='Active',data={'name':'Fictional Exam','campus':'Campus A','class_name':'Class A','papers':[{'subject':'Math','total':30,'date':'2026-09-01','type':'Written'}]})
        db.session.add(award);db.session.flush();db.session.add(self.ops['score'](award_id=award.id,student_id=self.s.id,paper=0,obtained=23.999,state=''));db.session.commit()
        self.post('exam','display',decimals='0',written='Theory')
        page=self.c.get('/workspace/core/exam/results/'+str(award.id));self.assertEqual(page.status_code,200);self.assertIn(b'Good',page.data);self.assertNotIn(b'Excellent',page.data)
        self.assertEqual(self.post('exam','grading',id=policy.id,version='1',action='delete').status_code,422)
    def test_weightage_and_test_types_are_separate_from_paper_total(self):
        self.assertEqual(self.post('exam','test-types',name='First Term',campus='Campus A').status_code,302);r=self.item('test-type')
        self.assertEqual(self.post('exam','weightage',campus='Campus A',**{'weight_'+str(r.id):'100.01'}).status_code,422)
        self.assertEqual(self.post('exam','weightage',campus='Campus A',**{'weight_'+str(r.id):'30'}).status_code,302)
        self.assertEqual(self.config('exam:weightage:Campus A').data[str(r.id)],'30')
        self.assertIn(b'value="" required',self.c.get('/workspace/examination/awards').data)
        self.assertEqual(self.post('exam','test-types',id=r.id,version='1',action='delete').status_code,422)
    def test_sms_saved_template_used_as_per_recipient_snapshot_no_delivery(self):
        self.assertEqual(self.post('hrm','sms',campus='Campus A',sms_0='Dear [emp_name], [date]: present.').status_code,302)
        self.assertEqual(self.c.post('/workspace/attendance/staff',data={'_csrf':self.csrf,'day':'2026-09-07','status_1':'Present','send':'1'}).status_code,302)
        message=self.ops['workflow'].query.filter_by(kind='message').one();self.assertEqual(message.status,'Provider not configured')
        self.assertIn('Dear Fictional Teacher, 2026-09-07: present.',message.data['messages'][0]['text'])
        self.assertEqual(self.post('hrm','sms',campus='Campus A',version='1',sms_0='[inventedTag]').status_code,422)
    def test_branding_sanitized_preview_no_writes_and_result_applies_campus(self):
        html='<h2>[institute_name]</h2><img src="https://evil.invalid/x" onerror="alert(1)"><script>alert(1)</script>'
        self.assertEqual(self.post('hrm','branding',name='Combine Attendance Report',campus='Campus A',header=html,body='[campus_name]',footer='Approved',action='preview').status_code,200)
        self.assertEqual(self.m['item'].query.count(),0)
        self.assertEqual(self.post('hrm','branding',name='Combine Attendance Report',campus='Campus A',header=html,body='[campus_name]',footer='Approved').status_code,302)
        stored=self.item('hrm-branding').data['header'];self.assertNotIn('<script',stored);self.assertNotIn('onerror',stored);self.assertNotIn('src=',stored)
        page=self.c.get('/workspace/reports/attendance/staff-attendance-report?campus=Campus%20A');self.assertEqual(page.status_code,200);self.assertIn(b'Approved',page.data)
        self.assertNotIn(b'Approved',self.c.get('/workspace/reports/attendance/staff-attendance-report?campus=Campus%20B').data)
    def test_employee_template_actual_record_escape_and_campus_scope(self):
        self.assertEqual(self.post('hrm','print-templates',name='Employment letter',campus='Campus A',body='<p>[emp_name] / [designation_name]</p>').status_code,302)
        t=self.item('print-template');page=self.c.get(f'/workspace/core/hrm/print-template/{t.id}?employee_id=1');self.assertEqual(page.status_code,200);self.assertIn(b'Fictional Teacher / Class Teacher',page.data)
        detail=self.rm['employee'].query.filter_by(employee_id=1).one();detail.data={'campus':'Campus B'};db.session.commit()
        self.assertEqual(self.c.get(f'/workspace/core/hrm/print-template/{t.id}?employee_id=1').status_code,404)
    def test_catalog_in_use_protected_and_designation_group_validated(self):
        r=self.rm['settings'].query.filter_by(category='designation',value='Class Teacher').one()
        self.assertEqual(self.post('hrm','designations',id=r.id,name='Other Teacher',group='Teaching').status_code,422)
        self.assertEqual(self.post('hrm','designations',id=r.id,name='Class Teacher',group='Teaching').status_code,302)
        with self.assertRaises(ValueError):app.extensions['school_settings_helpers']['employee_validate']({'role':'Class Teacher','designation_group':'Office'})
    def test_allowance_snapshot_used_in_payroll_and_not_recalculated(self):
        self.assertEqual(self.post('hrm','allowances',name='Transport',type='Earning',basic='on').status_code,302);head=self.item('allowance')
        plan=dict(_csrf=self.csrf,employee_id='1',basic='10000',**{'allowance_'+str(head.id):'1200'})
        for i in range(6):plan.update({f'allowed_{i}':'0',f'jump_{i}':'0',f'limit_{i}':'0'})
        self.assertEqual(self.c.post('/workspace/salary-plans',data=plan).status_code,302);row=self.ops['workflow'].query.filter_by(kind='salary-plans').one();row.status='Approved';db.session.commit()
        self.assertEqual(self.c.post('/workspace/payroll',data={'_csrf':self.csrf,'month':'2026-09','action':'generate','employee_ids':'1'}).status_code,302)
        salary=self.ops['salary'].query.one();self.assertEqual(salary.net,1120000)
        self.assertEqual(self.post('hrm','allowances',id=head.id,version='1',action='delete').status_code,422)
        self.assertEqual(salary.detail['plan']['allowances'][0]['name'],'Transport')

    def test_weighted_results_require_complete_config_and_do_not_normalize_missing_types(self):
        db.session.add(self.rm['settings'](category='academic_session',value='2026-27',class_name=''))
        detail=self.rm['student'].query.filter_by(student_id=self.s.id).one();detail.data={**detail.data,'academic_session':'2026-27'};db.session.commit()
        self.post('exam','test-types',name='Term 1',campus='Campus A');self.post('exam','test-types',name='Term 2',campus='Campus A')
        types=self.m['item'].query.filter_by(kind='test-type').order_by(self.m['item'].name).all()
        weights={'weight_'+str(types[0].id):'40','weight_'+str(types[1].id):'60'}
        self.post('exam','weightage',campus='Campus A',**weights)
        path='/workspace/core/exam/weighted-results?campus=Campus+A&class_name=Class+A&academic_session=2026-27'
        self.assertIn(b'Incomplete',self.c.get(path).data)
        for typ,total,obtained in [('Term 1',20,10),('Term 2',30,30)]:
            a=self.ops['workflow'](kind='awards',status='Active',data={'name':typ,'exam_test_type':typ,'campus':'Campus A','class_name':'Class A','academic_session':'2026-27','papers':[{'subject':'Math','total':total}]})
            db.session.add(a);db.session.flush();db.session.add(self.ops['score'](award_id=a.id,student_id=self.s.id,paper=0,obtained=obtained,state=''))
        db.session.commit();page=self.c.get(path);self.assertEqual(page.status_code,200);self.assertIn(b'<td>80.00</td>',page.data)
        self.post('exam','weightage',campus='Campus A',version='1',**{'weight_'+str(types[0].id):'40'})
        self.assertEqual(self.c.get(path).status_code,422)

    def test_shift_copy_is_same_designation_and_campus_only(self):
        self.shift();shift=self.item('shift-template')
        extra=Employee(employee_no='E-3',name='Other campus teacher',role='Class Teacher',status='Active');db.session.add(extra);db.session.flush();db.session.add(self.rm['employee'](employee_id=extra.id,data={'campus':'Campus B'}));db.session.commit()
        path='/workspace/core/hrm/assign-shifts'
        self.c.post(path,data={'_csrf':self.csrf,'version':'0','action':'save','employee_ids':'1','in_1_0':str(shift.id)})
        self.assertEqual(self.c.post(path,data={'_csrf':self.csrf,'version':'1','action':'copy','source':'1'}).status_code,302)
        result=self.config('hrm:shift-assignment').data;self.assertNotIn('2',result);self.assertNotIn(str(extra.id),result)

    def test_invalid_later_monthly_status_rolls_back_earlier_days(self):
        path='/workspace/core/hrm/monthly-attendance'
        self.assertEqual(self.c.post(path,data={'_csrf':self.csrf,'employee_id':'1','month':'2026-09','status_2026-09-01':'Present','status_2026-09-02':'Invalid'}).status_code,422)
        self.assertEqual(self.ops['daily'].query.count(),0)

    def test_required_uploads_validate_existing_record_and_removal(self):
        fields={}
        for _,fs in __import__('record_schema').EMPLOYEE_SECTIONS:
            for f in fs:fields['visible_'+f['key']]='1'
        fields.update(visible_photo='1',required_photo='1',visible_documents='1')
        self.assertEqual(self.post('hrm','employee-fields',**fields).status_code,302)
        form=dict(_csrf=self.csrf,name='No Photo',father_name='Guardian',gender='Female',marital_status='Married',phone='00000000000',campus='Campus A',department='Academics',designation_group='Teaching',role='Class Teacher',employment_type='Permanent',employee_no='E-3',joining_date='2000-01-01')
        page=self.c.post('/records/employee/new',data=form);self.assertEqual(page.status_code,422);self.assertIn(b'This upload is required',page.data);self.assertEqual(Employee.query.count(),2)

    def test_settings_excel_export_is_formula_safe_and_read_only(self):
        import io
        from openpyxl import load_workbook
        self.post('hrm','departments',name='=TEST-FORMULA')
        before=self.ops['audit'].query.count()
        response=self.c.get('/workspace/core/hrm/departments?export=xlsx')
        self.assertEqual(response.status_code,200)
        sheet=load_workbook(io.BytesIO(response.data)).active
        values=[row[0].value for row in sheet.iter_rows(min_row=2)]
        self.assertIn("'=TEST-FORMULA",values)
        self.assertEqual(self.ops['audit'].query.count(),before)

    def test_root_settings_forms_submit_to_their_visible_page(self):
        response=self.c.post('/workspace/core/exam',data={'_csrf':self.csrf,'version':'0','written':'Root Theory','decimals':'2'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.config('exam:display').data['written'],'Root Theory')
        response=self.c.post('/workspace/core/hrm',data={'_csrf':self.csrf,'version':'0','action':'reset'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.config('hrm:employee-fields').data,{})

    def test_salary_and_increment_prints_apply_the_selected_campus_branding(self):
        salary=self.ops['salary'](employee_id=1,month='2026-09',earnings=10000,deductions=0,net=10000,paid=0,detail={'basic':10000,'lecture_pay':0,'attendance':[]})
        increment=self.ops['workflow'](kind='increments',status='Approved',data={'employee_id':'1','amount':'100','effective':'2026-09','reason':'Test only'})
        db.session.add_all([salary,increment]);db.session.commit()
        self.post('hrm','branding',name='Bank Salary Sheet',campus='Campus A',footer='Campus A bank footer')
        page=self.c.post('/workspace/payroll',data={'_csrf':self.csrf,'month':'2026-09','action':'bank','ids':str(salary.id)})
        self.assertEqual(page.status_code,200);self.assertIn(b'Campus A bank footer',page.data)
        self.post('hrm','branding',name='Salary Increment Sheet',campus='Campus A',footer='Campus A increment footer')
        page=self.c.get('/workspace/core/hrm/increment-sheet?campus=Campus+A&from=2026-09&to=2026-09')
        self.assertEqual(page.status_code,200);self.assertIn(b'Campus A increment footer',page.data);self.assertIn(b'Test only',page.data)
        self.assertNotIn(b'Campus A increment footer',self.c.get('/workspace/core/hrm/increment-sheet?campus=Campus+B').data)

    def test_subject_reordering_is_versioned_atomic_and_used_in_admission_options(self):
        self.post('exam','subject-groups',name='Academics')
        self.post('exam','subjects',name='Zulu',group='Academics')
        self.post('exam','subjects',name='Alpha',group='Academics')
        rows=self.m['item'].query.filter_by(kind='subject').all();byname={r.name:r for r in rows}
        d=self.rm['student'].query.filter_by(student_id=1).one();d.data={**d.data,'subjects':['Zulu']}
        award=self.ops['workflow'](kind='awards',status='Active',data={'name':'Preserved award','papers':[{'subject':'Zulu','total':20}]})
        db.session.add(award);db.session.flush();score=self.ops['score'](award_id=award.id,student_id=1,paper=0,obtained=18);db.session.add(score);db.session.commit()
        ordered=[str(byname['Alpha'].id),str(byname['Zulu'].id)]
        versions={'version_'+str(r.id):str(r.version) for r in rows}
        self.assertEqual(self.post('exam','subjects',action='reorder',order_ids=ordered,**versions).status_code,302)
        page=self.c.get('/records/student/new').data
        subject_select=re.search(rb'<select id="subjects".*?</select>',page,re.S).group()
        self.assertLess(subject_select.index(b'Alpha'),subject_select.index(b'Zulu'))
        self.assertEqual(d.data['subjects'],['Zulu']);self.assertEqual(award.data['papers'][0]['subject'],'Zulu');self.assertEqual(score.obtained,18)
        self.assertEqual(self.post('exam','subjects',action='reorder',order_ids=list(reversed(ordered)),**versions).status_code,422)
        fresh={'version_'+str(r.id):str(r.version) for r in rows}
        self.assertEqual(self.post('exam','subjects',action='reorder',order_ids=[ordered[0],ordered[0]],**fresh).status_code,422)
        self.assertEqual([r.name for r in sorted(rows,key=lambda r:r.data['order'])],['Alpha','Zulu'])
        groups=self.m['item'].query.filter_by(kind='subject-group').all()
        self.assertIn(b'data-settings-order',self.c.get('/workspace/core/exam/subject-groups').data)
        self.assertEqual(self.post('exam','subject-groups',action='reorder',order_ids=[str(r.id) for r in groups],**{'version_'+str(r.id):str(r.version) for r in groups}).status_code,302)

    def test_legacy_class_subject_assignment_prevents_catalog_delete(self):
        self.post('exam','subject-groups',name='Academics');self.post('exam','subjects',name='Math',group='Academics')
        subject=self.item('subject');db.session.add(self.rm['settings'](category='subject',value='Math',class_name='Class A'));db.session.commit()
        self.assertEqual(self.post('exam','subjects',id=subject.id,version=str(subject.version),action='delete').status_code,422)
        self.assertEqual(self.rm['settings'].query.filter_by(category='subject',value='Math').count(),2)

    def test_editor_css_is_preserved_and_html_preview_cannot_execute_scripts(self):
        html='<p style="font-family:Georgia;font-size:18px;color:#176b4a;background-color:#e7f3ec" onclick="bad()">Safe heading</p><img src=x onerror="bad()">'
        response=self.post('hrm','branding',name='Salary Slip Report',header=html,action='preview')
        self.assertEqual(response.status_code,200);self.assertNotIn(b'onerror=',response.data);self.assertNotIn(b'onclick="bad()',response.data)
        self.assertIn(b'font-family:Georgia',response.data);self.assertIn(b'aria-label="header font"',response.data)
        self.assertEqual(self.m['item'].query.count(),0)

    def test_required_employee_checkbox_is_enforced(self):
        fields={'visible_app_active':'1','required_app_active':'1'}
        self.assertEqual(self.post('hrm','employee-fields',**fields).status_code,302)
        form=dict(_csrf=self.csrf,name='Required checkbox',father_name='Guardian',gender='Female',marital_status='Married',phone='00000000000',campus='Campus A',department='Academics',designation_group='Teaching',role='Class Teacher',employment_type='Permanent',employee_no='E-3',joining_date='2000-01-01')
        response=self.c.post('/records/employee/new',data=form);self.assertEqual(response.status_code,422);self.assertIn(b'This field must be checked.',response.data)
        form['app_active']='1';self.assertEqual(self.c.post('/records/employee/new',data=form).status_code,302)
